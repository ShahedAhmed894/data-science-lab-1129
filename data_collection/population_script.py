import requests
from bs4 import BeautifulSoup
import csv
import time
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# CSE 4112 - DATA SCIENCE LAB
# ============================================================
#
# Domain:
# Population / Demography
#
# Website:
# Worldometer
#
# Scraping:
# Requests + BeautifulSoup
#
# Output:
#   population_data.csv
#   population_data.xlsx
#
# Requirement:
#   4000+ records
# ============================================================


BASE_URL = "https://www.worldometers.info"

START_URL = (
        BASE_URL + "/world-population/"
)

CSV_FILE = "population_data.csv"

EXCEL_FILE = "population_data.xlsx"

TARGET_RECORDS = 4000

DELAY = 0.5


# ============================================================
# HTTP SESSION
# ============================================================

session = requests.Session()

session.headers.update({

    "User-Agent":
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36",

    "Accept":
        "text/html,application/xhtml+xml,"
        "application/xml;q=0.9,image/avif,image/webp,"
        "*/*;q=0.8",

    "Accept-Language":
        "en-US,en;q=0.9"

})


# ============================================================
# CLEAN TEXT
# ============================================================

def clean_text(value):

    if value is None:
        return ""

    value = str(value)

    # Fix common encoding issues
    replacements = {

        "â": "-",
        "â€“": "-",
        "â€”": "-",

        "−": "-",
        "–": "-",
        "—": "-",

        "\xa0": " "
    }

    for old, new in replacements.items():

        value = value.replace(
            old,
            new
        )

    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip()


# ============================================================
# CLEAN NUMERIC DATA
# ============================================================

def clean_numeric(value):

    value = clean_text(
        value
    )

    if not value:

        return None


    # Remove comma separators
    value = value.replace(
        ",",
        ""
    )


    # Remove percentage
    value = value.replace(
        "%",
        ""
    )


    # Remove currency symbols
    value = value.replace(
        "$",
        ""
    )


    # Handle missing values
    if value.lower() in [
        "",
        "n/a",
        "na",
        "null",
        "none",
        "-"
    ]:

        return None


    try:

        number = float(
            value
        )

        if number.is_integer():

            return int(
                number
            )

        return number

    except ValueError:

        return None


# ============================================================
# GET SOUP
# ============================================================

def get_soup(url):

    try:

        response = session.get(
            url,
            timeout=30
        )

        print(
            f"HTTP {response.status_code}: {url}"
        )

        response.raise_for_status()

        return BeautifulSoup(
            response.text,
            "html.parser"
        )

    except Exception as error:

        print(
            "ERROR:",
            error
        )

        return None


# ============================================================
# DISCOVER COUNTRY PAGES
# ============================================================

def discover_country_pages():

    print()
    print("=" * 70)
    print("DISCOVERING COUNTRY PAGES")
    print("=" * 70)


    soup = get_soup(
        START_URL
    )


    if soup is None:

        return []


    pages = {}


    for link in soup.find_all(
            "a",
            href=True
    ):

        href = link["href"]


        if not href.startswith(
                "/world-population/"
        ):

            continue


        if not href.endswith(
                "-population/"
        ):

            continue


        if href == "/world-population/":

            continue


        url = BASE_URL + href


        country = clean_text(
            link.get_text(
                " ",
                strip=True
            )
        )


        if not country:

            slug = href.strip(
                "/"
            ).split("/")[-1]

            slug = slug.replace(
                "-population",
                ""
            )

            country = slug.replace(
                "-",
                " "
            ).title()


        pages[url] = country


    result = list(
        pages.items()
    )


    print(
        "Country pages discovered:",
        len(result)
    )


    return result


# ============================================================
# FIND POPULATION TABLE
# ============================================================

def find_population_table(soup):

    tables = soup.find_all(
        "table"
    )


    print(
        "HTML tables found:",
        len(tables)
    )


    for table in tables:

        rows = table.find_all(
            "tr"
        )


        if len(rows) < 10:

            continue


        # ----------------------------------------------------
        # Inspect rows for Year + Population
        # ----------------------------------------------------

        for row in rows[:10]:

            cells = row.find_all(
                ["th", "td"]
            )


            headers = [
                clean_text(
                    cell.get_text(
                        " ",
                        strip=True
                    )
                ).lower()
                for cell in cells
            ]


            if (
                    "year" in headers
                    and
                    "population" in headers
            ):

                return table


    return None


# ============================================================
# NORMALIZE HEADER
# ============================================================

def normalize_header(header):

    header = clean_text(
        header
    )

    lower = header.lower()


    # --------------------------------------------------------
    # YEAR
    # --------------------------------------------------------

    if lower == "year":

        return "Year"


    # --------------------------------------------------------
    # POPULATION
    # --------------------------------------------------------

    if lower == "population":

        return "Population"


    # --------------------------------------------------------
    # YEARLY PERCENT CHANGE
    # --------------------------------------------------------

    if (
            "yearly" in lower
            and
            "%" in lower
    ):

        return "Yearly_Percent_Change"


    if (
            "yearly" in lower
            and
            "change" in lower
            and
            "%" not in lower
    ):

        return "Yearly_Change"


    # --------------------------------------------------------
    # MIGRANTS
    # --------------------------------------------------------

    if "migrant" in lower:

        return "Migrants_Net"


    # --------------------------------------------------------
    # MEDIAN AGE
    # --------------------------------------------------------

    if "median age" in lower:

        return "Median_Age"


    # --------------------------------------------------------
    # FERTILITY
    # --------------------------------------------------------

    if "fertility" in lower:

        return "Fertility_Rate"


    # --------------------------------------------------------
    # DENSITY
    # --------------------------------------------------------

    if "density" in lower:

        return "Density_Per_Km2"


    # --------------------------------------------------------
    # URBAN POPULATION PERCENT
    # --------------------------------------------------------

    if (
            "urban" in lower
            and
            "%" in lower
    ):

        return "Urban_Population_Percent"


    # --------------------------------------------------------
    # URBAN POPULATION
    # --------------------------------------------------------

    if (
            "urban" in lower
            and
            "population" in lower
    ):

        return "Urban_Population"


    # --------------------------------------------------------
    # COUNTRY SHARE
    # --------------------------------------------------------

    if (
            "share" in lower
            and
            "world" in lower
    ):

        return (
            "Country_Share_of_World_Population"
        )


    # --------------------------------------------------------
    # WORLD POPULATION
    # --------------------------------------------------------

    if (
            "world" in lower
            and
            "population" in lower
    ):

        return "World_Population"


    # --------------------------------------------------------
    # GLOBAL RANK
    # --------------------------------------------------------

    if (
            "global" in lower
            and
            "rank" in lower
    ):

        return "Global_Rank"


    # --------------------------------------------------------
    # FALLBACK
    # --------------------------------------------------------

    result = re.sub(
        r"[^A-Za-z0-9]+",
        "_",
        header
    )

    result = result.strip(
        "_"
    )


    return result


# ============================================================
# PARSE TABLE
# ============================================================

def parse_table(
        table,
        country
):

    rows = table.find_all(
        "tr"
    )


    # --------------------------------------------------------
    # Find header row
    # --------------------------------------------------------

    header_index = None

    raw_headers = []


    for index, row in enumerate(
            rows[:15]
    ):

        cells = row.find_all(
            ["th", "td"]
        )


        headers = [
            clean_text(
                cell.get_text(
                    " ",
                    strip=True
                )
            )
            for cell in cells
        ]


        lower = [
            header.lower()
            for header in headers
        ]


        if (
                "year" in lower
                and
                "population" in lower
        ):

            header_index = index

            raw_headers = headers

            break


    if header_index is None:

        print(
            "Header row not found."
        )

        return []


    # --------------------------------------------------------
    # Normalize headers
    # --------------------------------------------------------

    headers = [
        normalize_header(
            header
        )
        for header in raw_headers
    ]


    print(
        "Detected columns:"
    )


    for header in headers:

        print(
            "  ",
            header
        )


    # --------------------------------------------------------
    # Check required fields
    # --------------------------------------------------------

    if "Year" not in headers:

        print(
            "Year column missing."
        )

        return []


    if "Population" not in headers:

        print(
            "Population column missing."
        )

        return []


    records = []


    # ========================================================
    # DATA ROWS
    # ========================================================

    for row in rows[
        header_index + 1:
    ]:

        cells = row.find_all(
            ["td", "th"]
        )


        values = [
            clean_text(
                cell.get_text(
                    " ",
                    strip=True
                )
            )
            for cell in cells
        ]


        # Ignore malformed rows
        if len(values) != len(headers):

            continue


        # ----------------------------------------------------
        # Year
        # ----------------------------------------------------

        year_value = clean_numeric(
            values[
                headers.index("Year")
            ]
        )


        if year_value is None:

            continue


        if not (
                1900 <= year_value <= 2100
        ):

            continue


        # ----------------------------------------------------
        # Build record
        # ----------------------------------------------------

        record = {

            "Country": country,

            "Year": int(
                year_value
            )
        }


        for index, header in enumerate(
                headers
        ):

            if header == "Year":

                continue


            value = clean_numeric(
                values[index]
            )


            record[header] = value


        records.append(
            record
        )


    return records


# ============================================================
# SCRAPE COUNTRY
# ============================================================

def scrape_country(
        country,
        url
):

    print()
    print("=" * 70)
    print(
        f"SCRAPING: {country}"
    )
    print(
        url
    )
    print("=" * 70)


    soup = get_soup(
        url
    )


    if soup is None:

        return []


    table = find_population_table(
        soup
    )


    if table is None:

        print(
            "Population table NOT found."
        )

        return []


    records = parse_table(
        table,
        country
    )


    print(
        "Records collected:",
        len(records)
    )


    return records


# ============================================================
# REMOVE DUPLICATES
# ============================================================

def remove_duplicates(
        records
):

    unique = []

    seen = set()


    for record in records:

        key = (
            record["Country"],
            record["Year"]
        )


        if key in seen:

            continue


        seen.add(
            key
        )

        unique.append(
            record
        )


    return unique


# ============================================================
# CSV COLUMN ORDER
# ============================================================

COLUMNS = [

    "Country",

    "Year",

    "Population",

    "Yearly_Percent_Change",

    "Yearly_Change",

    "Migrants_Net",

    "Median_Age",

    "Fertility_Rate",

    "Density_Per_Km2",

    "Urban_Population_Percent",

    "Urban_Population",

    "Country_Share_of_World_Population",

    "World_Population",

    "Global_Rank"
]


# ============================================================
# SAVE CSV
# ============================================================

def save_csv(records):

    print()
    print("=" * 70)
    print("CREATING CSV")
    print("=" * 70)


    with open(
            CSV_FILE,
            "w",
            newline="",
            encoding="utf-8-sig"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=COLUMNS,
            extrasaction="ignore"
        )


        writer.writeheader()


        for record in records:

            writer.writerow({

                column:
                    record.get(
                        column,
                        ""
                    )

                for column in COLUMNS

            })


    print(
        f"CSV saved: {CSV_FILE}"
    )


# ============================================================
# SAVE EXCEL
# ============================================================

def save_excel(records):

    print()
    print("=" * 70)
    print("CREATING EXCEL FILE")
    print("=" * 70)


    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Population Data"


    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    sheet.append(
        COLUMNS
    )


    # --------------------------------------------------------
    # Header formatting
    # --------------------------------------------------------

    for cell in sheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )


    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    for record in records:

        sheet.append([

            record.get(
                column,
                None
            )

            for column in COLUMNS

        ])


    # --------------------------------------------------------
    # Freeze header
    # --------------------------------------------------------

    sheet.freeze_panes = "A2"


    # --------------------------------------------------------
    # Auto filter
    # --------------------------------------------------------

    sheet.auto_filter.ref = (
        sheet.dimensions
    )


    # --------------------------------------------------------
    # Create Excel table
    # --------------------------------------------------------

    table = Table(
        displayName="PopulationData",
        ref=sheet.dimensions
    )


    style = TableStyleInfo(

        name="TableStyleMedium2",

        showFirstColumn=False,

        showLastColumn=False,

        showRowStripes=True,

        showColumnStripes=False

    )


    table.tableStyleInfo = style

    sheet.add_table(
        table
    )


    # --------------------------------------------------------
    # Number formats
    # --------------------------------------------------------

    for row in sheet.iter_rows(
            min_row=2
    ):

        for cell in row:

            # Population
            if cell.column in [
                3,
                5,
                6,
                11,
                13
            ]:

                cell.number_format = (
                    '#,##0'
                )


            # Percentage
            elif cell.column in [
                4,
                10,
                12
            ]:

                cell.number_format = (
                    '0.00'
                )


            # Decimal
            elif cell.column in [
                7,
                8
            ]:

                cell.number_format = (
                    '0.00'
                )


            # Density
            elif cell.column == 9:

                cell.number_format = (
                    '#,##0'
                )


            # Rank
            elif cell.column == 14:

                cell.number_format = (
                    '0'
                )


    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    widths = {

        "A": 22,
        "B": 10,
        "C": 18,
        "D": 24,
        "E": 18,
        "F": 16,
        "G": 14,
        "H": 16,
        "I": 18,
        "J": 26,
        "K": 20,
        "L": 32,
        "M": 20,
        "N": 14
    }


    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width


    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    workbook.save(
        EXCEL_FILE
    )


    print(
        f"Excel saved: {EXCEL_FILE}"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 70)
    print("CSE 4112 - DATA SCIENCE LAB")
    print("GLOBAL POPULATION WEB SCRAPER")
    print("=" * 70)


    # --------------------------------------------------------
    # Discover countries
    # --------------------------------------------------------

    countries = discover_country_pages()


    if not countries:

        print(
            "No country pages found."
        )

        return


    all_records = []


    # --------------------------------------------------------
    # Scrape
    # --------------------------------------------------------

    for index, (
            url,
            country
    ) in enumerate(
        countries,
        start=1
    ):

        print()
        print(
            f"COUNTRY {index}/{len(countries)}"
        )


        records = scrape_country(
            country,
            url
        )


        all_records.extend(
            records
        )


        print(
            "Total raw records:",
            len(all_records)
        )


        # ----------------------------------------------------
        # Stop after safely exceeding requirement
        # ----------------------------------------------------

        if len(all_records) >= 4500:

            print()
            print(
                "4000+ records reached."
            )

            break


        time.sleep(
            DELAY
        )


    # --------------------------------------------------------
    # Remove duplicates
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("REMOVING DUPLICATES")
    print("=" * 70)


    raw_count = len(
        all_records
    )


    all_records = remove_duplicates(
        all_records
    )


    print(
        "Raw records:",
        raw_count
    )


    print(
        "Unique records:",
        len(all_records)
    )


    # --------------------------------------------------------
    # Sort
    # --------------------------------------------------------

    all_records.sort(
        key=lambda x: (
            x["Country"],
            x["Year"]
        )
    )


    # --------------------------------------------------------
    # Save CSV
    # --------------------------------------------------------

    save_csv(
        all_records
    )


    # --------------------------------------------------------
    # Save Excel
    # --------------------------------------------------------

    save_excel(
        all_records
    )


    # --------------------------------------------------------
    # Final result
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("FINAL RESULT")
    print("=" * 70)


    print(
        f"Total records: {len(all_records)}"
    )


    print(
        f"CSV: {CSV_FILE}"
    )


    print(
        f"Excel: {EXCEL_FILE}"
    )


    if len(all_records) >= 4000:

        print()
        print(
            "✓ REQUIREMENT PASSED"
        )

        print(
            "✓ 4000+ records"
        )

        print(
            "✓ Multiple HTML pages"
        )

        print(
            "✓ BeautifulSoup"
        )

        print(
            "✓ Organized CSV"
        )

        print(
            "✓ Excel version created"
        )

    else:

        print()
        print(
            "⚠ Fewer than 4000 records."
        )


    print()
    print(
        "Columns:"
    )


    for column in COLUMNS:

        print(
            " -",
            column
        )


    print()
    print("=" * 70)
    print("SCRAPING COMPLETED")
    print("=" * 70)


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    main()